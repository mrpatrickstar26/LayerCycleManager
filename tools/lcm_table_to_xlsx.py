# -*- coding: utf-8 -*-
"""Добавляет в «Уч. 6 шпунт.xlsx» второй блок таблицы вплотную справа (E:H).

Логика (по договорённости с пользователем):
  * копируется вся первая таблица (шапка + все строки марок, включая I1…I24);
  * значения отклонений сначала не переносятся;
  * в шапке номер цикла увеличивается на 1, дата — сегодняшняя;
  * отклонения подставляются из таблицы AutoCAD (reports/_table_cells.json)
    по совпадению марки (числовые и вида I16);
  * где в AutoCAD прочерк или марки нет — в Excel тоже прочерк.

Запуск:
    python tools/lcm_table_to_xlsx.py
"""
import json
import os
import re
import sys
from copy import copy
from datetime import date

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "Уч. 6 шпунт.xlsx")
DATA = os.path.join(ROOT, "reports", "_table_cells.json")

START_COL = 5          # второй блок: E:H, вплотную к первому (A:D)


def style_from(src, dst):
    dst.font = copy(src.font)
    dst.border = copy(src.border)
    dst.fill = copy(src.fill)
    dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format


def to_number(text):
    s = str(text).strip().replace(",", ".")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def mark_key(value):
    """Ключ марки: 57 -> 57, 'I16' -> 'I16' (без регистра)."""
    s = str(value).strip()
    return int(s) if s.isdigit() else s.upper()


def next_cycle(label):
    """«Цикл 0 08.10.2026» -> «Цикл 1 30.10.2026»."""
    today = date.today().strftime("%d.%m.%Y")
    parts = str(label).split()
    if len(parts) >= 2 and parts[1].isdigit():
        return f"{parts[0]} {int(parts[1]) + 1} {today}"
    return f"{label} {today}"


def main():
    with open(DATA, encoding="utf-8-sig") as fh:
        raw = json.load(fh)
    rows = [r["value"] if isinstance(r, dict) and "value" in r else r for r in raw]

    drawn = {}
    for row in rows[3:]:
        dx = to_number(row[1])
        if dx is not None:
            drawn[mark_key(row[0])] = (dx, to_number(row[2]), to_number(row[3]))

    wb = load_workbook(XLSX)
    ws = wb.active
    # Столбец D содержит формулы — просим Excel пересчитать их при открытии
    wb.calculation.fullCalcOnLoad = True

    c0 = START_COL
    col = lambda i: get_column_letter(c0 + i)

    # ---- убрать старый второй блок, если он был ----
    for rng in (f"{col(0)}1:{col(0)}3", f"{col(1)}1:{col(3)}1", f"{col(1)}2:{col(3)}2"):
        if rng in [str(m) for m in ws.merged_cells.ranges]:
            ws.unmerge_cells(rng)
    for r in range(1, ws.max_row + 1):
        for i in range(4):
            ws.cell(r, c0 + i).value = None

    # ---- шапка ----
    for i, src_col in enumerate(range(1, 5)):
        for r in (1, 2, 3):
            style_from(ws.cell(r, src_col), ws.cell(r, c0 + i))

    ws[f"{col(0)}1"] = ws["A1"].value        # «№ марки»
    ws[f"{col(1)}1"] = ws["B1"].value        # «Смещение по x и y (мм)»
    ws[f"{col(1)}2"] = next_cycle(ws["B2"].value)
    ws[f"{col(1)}3"] = ws["B3"].value        # «ΣΔX»
    ws[f"{col(2)}3"] = ws["C3"].value        # «ΣΔY»
    ws[f"{col(3)}3"] = ws["D3"].value        # «D»

    ws.merge_cells(f"{col(0)}1:{col(0)}3")
    ws.merge_cells(f"{col(1)}1:{col(3)}1")
    ws.merge_cells(f"{col(1)}2:{col(3)}2")

    # ---- строки марок: ровно как в первом блоке ----
    filled = 0
    for r in range(4, ws.max_row + 1):
        mark = ws.cell(r, 1).value
        if mark is None:                      # пустая строка-разделитель
            continue

        style_from(ws.cell(r, 1), ws.cell(r, c0))
        ws.cell(r, c0).value = mark

        values = drawn.get(mark_key(mark))
        for offset in (1, 2, 3):
            cell = ws.cell(r, c0 + offset)
            style_from(ws.cell(r, 1 + offset), cell)

            if values is None:
                cell.value = "-"
                cell.number_format = "0"
                continue

            num = values[offset - 1]
            if offset == 3:
                # D считается по той же формуле, что и в первом блоке
                cell.value = f"=({col(1)}{r}*{col(1)}{r}+{col(2)}{r}*{col(2)}{r})^0.5"
                cell.number_format = "0"
            else:
                cell.value = num
                cell.number_format = "0.0"
        if values is not None:
            filled += 1

    # ширина второго блока = ширине первого
    for i in range(4):
        ws.column_dimensions[col(i)].width = ws.column_dimensions[get_column_letter(1 + i)].width

    wb.save(XLSX)
    print(f"готово: {XLSX}; блок {col(0)}:{col(3)}; "
          f"цикл: {ws[f'{col(1)}2'].value}; марок с отклонениями: {filled}")


if __name__ == "__main__":
    sys.exit(main())
